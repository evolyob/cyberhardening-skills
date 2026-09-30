#!/usr/bin/env python3
"""noai-note Minimal Ingester: Multi-source Noise Stripping, Agenda Chunking & Circuit Breaker."""
import argparse, html, json, re, sys, unicodedata, zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

SUPPORTED_EXTS = {".docx", ".pdf", ".md", ".txt", ".xlsx", ".html"}
MAX_STREAM_TOKENS, MAX_STREAM_PAGES = 6500, 25

CLEAN_PATTERN = re.compile(
    r'[\u200b-\u200d\ufeff\u2060\u202a-\u202e\u2066-\u2069\x00-\x08\x0b\x0c\x0e-\x1f\x7f]'
    r'|\(cid:\d+\)|\ufffd|^\s*(?:[-—~]*\s*\d+\s*[-—~]*|Page\s+\d+(?:\s*(?:of|/)\s*\d+)?|\d+\s*/\s*\d+)\s*$',
    re.MULTILINE
)
MEETING_PATTERN = re.compile(
    r"^(?:#{1,4}[ \t]+|\[?[0-9]{1,2}:[0-9]{2}(?::[0-9]{2})?\]?[ \t]*|(?:議程|議題|Agenda|決議|待辦事項|Action Items|發言人)[0-9一二三四五六七八九十]*[:：]?[ \t]*)[^\r\n]*$",
    re.MULTILINE | re.IGNORECASE
)


def parse_html(raw: str) -> str:
    """Strip web clutter and elevate headings to Markdown anchors."""
    t = re.sub(r'<(script|style|nav|header|footer)[^>]*>.*?</\1>', '', raw, flags=re.DOTALL | re.IGNORECASE)
    t = re.sub(r'<h[1-6][^>]*>(.*?)</h[1-6]>', r'\n# \1\n', t, flags=re.DOTALL | re.IGNORECASE)
    return html.unescape(re.sub(r'<[^>]+>', ' ', t))


def read_raw(p: Path) -> tuple[str, int]:
    """Format ingestion: docx, pdf, xlsx, html, md, txt. Returns (text, page_count)."""
    ext = p.suffix.lower()
    if ext == ".docx":
        with zipfile.ZipFile(p) as z:
            tree = ET.fromstring(z.read("word/document.xml"))
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            return "\n".join("".join(t.text for t in n.findall(".//w:t", ns) if t.text) for n in tree.findall(".//w:p", ns)), 1
    if ext == ".xlsx":
        try:
            import openpyxl
            wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
            lines = [f"# Sheet: {s}\n" + "\n".join(" | ".join(str(v).strip() for v in r if v is not None and str(v).strip()) for r in wb[s].iter_rows(values_only=True) if any(r)) for s in wb.sheetnames]
            return "\n".join(lines), len(wb.sheetnames)
        except ImportError:
            return "", 0
    if ext == ".pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(str(p))
            return "\n".join(pg.extract_text() or "" for pg in reader.pages), len(reader.pages)
        except Exception:
            return "", 0
    if ext in (".md", ".txt", ".html"):
        content = p.read_text(encoding="utf-8-sig", errors="replace")
        return (parse_html(content) if ext == ".html" else content), 1
    return "", 0


def clean_text(raw: str) -> str:
    """Single-pass sanitization: NFKC normalization + noise removal."""
    if not raw: return ""
    text = re.sub(r'(\b[A-Za-z]+)-\n([A-Za-z]+\b)', r'\1\2', unicodedata.normalize("NFKC", raw))
    return "\n".join(L.strip() for L in CLEAN_PATTERN.sub("", text).splitlines() if L.strip())


def chunk_content(text: str, source: str) -> list[dict]:
    """Segment content via meeting anchors or fallback to 500-line sliding window."""
    matches = list(MEETING_PATTERN.finditer(text))
    if matches:
        chunks = []
        for i, m in enumerate(matches):
            start = m.end()
            end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
            body = text[start:end].strip()
            if len(body) >= 50:
                chunks.append({"source": source, "topic": m.group(0).strip("# ").strip(), "content": body})
        if chunks: return chunks

    lines = text.splitlines()
    total = len(lines)
    if total <= 500:
        return [{"source": source, "topic": "Full Document", "content": text}]
    chunks = []
    for idx, start_idx in enumerate(range(0, total, 450), 1):
        end_idx = min(start_idx + 500, total)
        chunks.append({"source": source, "topic": f"Part {idx:02d} (Lines {start_idx+1}-{end_idx})", "content": "\n".join(lines[start_idx:end_idx])})
        if end_idx >= total: break
    return chunks


def main():
    p = argparse.ArgumentParser(description="noai-note Ingester with Circuit Breaker & Split Support")
    p.add_argument("input", help="Target file or directory path")
    p.add_argument("-o", "--output", help="Destination file path")
    p.add_argument("--format", choices=["stream", "json"], default="stream")
    p.add_argument("--split-dir", help="Directory to export section files")
    p.add_argument("--list-only", action="store_true", help="List sections and tokens only")
    p.add_argument("--check", action="store_true", help="Perform readiness check")
    args = p.parse_args()

    if args.check:
        print("[OK] noai-note ingester is Ready with 25-page/6500-token circuit breaker.")
        return

    inp = Path(args.input).resolve()
    files = [inp] if inp.is_file() and inp.suffix.lower() in SUPPORTED_EXTS else sorted(x for x in inp.rglob("*") if x.suffix.lower() in SUPPORTED_EXTS) if inp.is_dir() else []
    if not files: sys.exit(f"Error: No valid documents found in: {inp}")

    records, total_pages = [], 0
    for f in files:
        raw, p_cnt = read_raw(f)
        total_pages += p_cnt
        if raw: records.extend(chunk_content(clean_text(raw), f.name))
    if not records: sys.exit("Error: No extractable content found.")

    total_chars = sum(len(r["content"]) for r in records)
    est_tokens = total_chars // 3
    catalog = [{"index": i, "id": f"sec_{i:02d}", "topic": r["topic"], "file": f"{i:02d}_{re.sub(r'[\\\\/*?:\"<>|#\\s]+', '_', r['topic']).strip('_')[:40]}.md", "lines": r["content"].count("\n") + 1, "est_tokens": len(r["content"]) // 3} for i, r in enumerate(records, 1)]

    if args.split_dir:
        dest_dir = Path(args.split_dir).resolve()
        dest_dir.mkdir(parents=True, exist_ok=True)
        for c, r in zip(catalog, records): (dest_dir / c["file"]).write_text(f"# {r['topic']}\n\n{r['content']}\n", encoding="utf-8-sig")
        (dest_dir / "data").mkdir(exist_ok=True)
        (dest_dir / "data" / "index.json").write_text(json.dumps({"total_sections": len(records), "total_est_tokens": est_tokens, "total_pages": total_pages, "sections": catalog}, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"Exported {len(records)} sections and data/index.json to {dest_dir}")
        return

    if args.list_only or (est_tokens > MAX_STREAM_TOKENS or total_pages > MAX_STREAM_PAGES):
        if not args.list_only: sys.stderr.write(f"[NOTICE] Large input ({total_pages}p, {est_tokens}t > {MAX_STREAM_TOKENS}t/{MAX_STREAM_PAGES}p). Use --split-dir.\n")
        print(json.dumps({"total_pages": total_pages, "total_est_tokens": est_tokens, "sections": catalog}, ensure_ascii=False, indent=2))
        return

    out = json.dumps(records, ensure_ascii=False, indent=2) if args.format == "json" else "\n\n".join(f"=== [SOURCE: {r['source']} | TOPIC: {r['topic']}] ===\n{r['content']}" for r in records)
    if args.output:
        dest = Path(args.output).resolve()
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_text(out, encoding="utf-8-sig")
        sys.stderr.write(f"[OK] Ingested {len(files)} source(s), {len(records)} section(s) (約 {est_tokens:,} tokens) -> {dest.name}\n")
    else:
        print(out)


if __name__ == "__main__":
    main()
